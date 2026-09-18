"""Plain-text extraction from client-provided documents (PDF, DOCX, XLSX, CSV).

Produces plain text only -- no chunking, no anonymization, no Qdrant. `ingestion.py` anonymizes
whatever this module returns before anything downstream (embedding, storage) ever sees it.

XLSX/CSV need one extra step DOCX/PDF don't: a spreadsheet's identity-bearing values (a name, an
email, a login) usually sit in their own cell with no keyword next to them -- "Marie Dupont" under
a column header "Responsable", not the sentence "Responsable : Marie Dupont" the anonymizer's
keyword rules are built to match (see src/anonymization/anonymizer.py). Each row is rendered as
"{column}: {cell}" pairs, and a cell whose column header identifies it as personal/identity data
is instead rendered with a canonical trigger keyword prepended ("Responsable : Marie Dupont"), so
the exact same anonymization rules used everywhere else in this project catch it too -- no
document-specific anonymization logic, just giving the existing rules something to trigger on.
"""

import csv
import re
import unicodedata
from pathlib import Path

SUPPORTED_EXTENSIONS = frozenset({".pdf", ".docx", ".xlsx", ".csv"})


class ExtractionError(RuntimeError):
    pass


# column header (normalised: lowercased, accents stripped) substring -> canonical keyword the
# anonymizer's rules already recognise (see default_rules() in anonymizer.py).
_IDENTITY_COLUMNS = {
    "responsable": "Responsable",
    "proprietaire": "Propriétaire",
    "owner": "Owner",
    "contact": "Contact",
    "nom": "Nom",
    "utilisateur": "Utilisateur",
    "username": "Username",
    "login": "Login",
    "email": "Email",
    "e-mail": "Email",
    "serveur": "Serveur",
    "hote": "Hôte",
    "host": "Host",
    "application": "Application",
    "societe": "Société",
    "entreprise": "Entreprise",
    "company": "Company",
}


def _normalize_header(header: str) -> str:
    ascii_header = unicodedata.normalize("NFKD", header).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "", ascii_header.lower())


def _identity_keyword(header: str) -> str | None:
    normalized = _normalize_header(header)
    for needle, keyword in _IDENTITY_COLUMNS.items():
        if needle in normalized:
            return keyword
    return None


def _row_to_text(headers: list, cells: list) -> str:
    pairs = []
    for header, cell in zip(headers, cells):
        value = "" if cell is None else str(cell).strip()
        if not value:
            continue
        keyword = _identity_keyword(header) if header else None
        pairs.append(f"{keyword} : {value}" if keyword else f"{header}: {value}")
    return ", ".join(pairs)


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError("pypdf is required to import PDF documents") from exc
    try:
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ExtractionError(f"Failed to read PDF '{path.name}': {exc}") from exc
    return "\n\n".join(p.strip() for p in pages if p.strip())


def extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ExtractionError("python-docx is required to import DOCX documents") from exc
    try:
        document = docx.Document(str(path))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))
    except Exception as exc:
        raise ExtractionError(f"Failed to read DOCX '{path.name}': {exc}") from exc
    return "\n\n".join(parts)


def extract_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise ExtractionError("openpyxl is required to import XLSX documents") from exc
    workbook = None
    try:
        workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        sheets_text = []
        for sheet in workbook.worksheets:
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [str(h).strip() if h is not None else "" for h in rows[0]]
            lines = [_row_to_text(headers, list(row)) for row in rows[1:]]
            lines = [line for line in lines if line]
            if lines:
                sheets_text.append(f"[{sheet.title}]\n" + "\n".join(lines))
    except Exception as exc:
        raise ExtractionError(f"Failed to read XLSX '{path.name}': {exc}") from exc
    finally:
        # A read-only workbook keeps its zip file handle open until closed explicitly; without
        # this, deleting the source file right after extraction fails on Windows (PermissionError)
        # -- found live, importing a real upload through a temp-directory cleanup.
        if workbook is not None:
            workbook.close()
    return "\n\n".join(sheets_text)


def extract_csv(path: Path) -> str:
    try:
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            rows = list(reader)
    except OSError as exc:
        raise ExtractionError(f"Failed to read CSV '{path.name}': {exc}") from exc
    if not rows:
        return ""
    headers = rows[0]
    lines = [_row_to_text(headers, row) for row in rows[1:]]
    return "\n".join(line for line in lines if line)


_EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
    ".csv": extract_csv,
}


def extract_text(path: Path) -> str:
    path = Path(path)
    extractor = _EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        raise ExtractionError(
            f"Unsupported document type '{path.suffix}' ({path.name}); supported: "
            f"{', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    return extractor(path)
