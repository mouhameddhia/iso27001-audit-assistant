"""Full auditor workflow through the real HTTP API, live stack: login -> create session -> import
a client document -> review its sidecar -> confirm -> draft a finding referencing the same
entities -> approve it -> download the report -> verify real values are restored and no
placeholder ever leaks. The API equivalent of tests/test_document_e2e.py and tests/test_report_e2e.py,
proving the web layer didn't regress any of the guarantees those already validate.
"""

import io
import re

import docx
import pytest

from src.config import get_settings
from tests.conftest import huggingface_reachable, ollama_generation_model_available, ollama_model_available
from tests.test_api.conftest import auth_headers, make_user

pytestmark = pytest.mark.e2e

PLACEHOLDER_TOKEN_RE = re.compile(r"\b[A-Z]+_\d{3}\b")


@pytest.fixture(autouse=True)
def _skip_if_unreachable():
    settings = get_settings()
    if not (ollama_model_available(settings) and ollama_generation_model_available(settings) and huggingface_reachable()):
        pytest.skip("Embedding/generation Ollama model or huggingface.co not reachable")


def test_full_workflow_through_the_api(client, db_session):
    make_user(db_session, username="jdupont")
    headers = auth_headers(client, "jdupont")

    r = client.post("/sessions", headers=headers, json={
        "title": "Rapport ISO/IEC 27001 - Meridian API", "client_name": "Meridian Holdings",
        "client_aliases": ["Meridian"], "scope": "Perimetre API e2e.",
        "standards": ["ISO/IEC 27001:2022"], "reference": "MISSION-API-E2E-01",
    })
    assert r.status_code == 201, r.text
    session_id = r.json()["id"]

    document = docx.Document()
    document.add_paragraph("Rapport d'audit precedent.")
    document.add_paragraph("Le present rapport concerne l'audit de securite realise pour Meridian en 2024.")
    document.add_paragraph(
        "Le serveur SRV-API-E2E-77 conservait des journaux d'audit non chiffres depuis sa mise "
        "en service sans que cela ait ete corrige a la suite de la precedente mission."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    r = client.post(
        f"/sessions/{session_id}/documents", headers=headers,
        files={"file": (
            "rapport_precedent.docx", buffer,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )},
        data={"doc_type": "previous_report"},
    )
    assert r.status_code == 201, r.text
    doc_id = r.json()["doc_id"]

    r = client.get(f"/sessions/{session_id}/documents/{doc_id}/sidecar", headers=headers)
    assert "Meridian" not in r.json()["content"]
    assert "SRV-API-E2E-77" not in r.json()["content"]

    r = client.post(f"/sessions/{session_id}/documents/{doc_id}/confirm", headers=headers)
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "confirmed"

    observation = (
        "Chez Meridian, le serveur SRV-API-E2E-77 conserve encore des journaux d'audit non "
        "chiffres, comme deja releve dans un rapport precedent."
    )
    r = client.post(f"/sessions/{session_id}/findings", headers=headers, json={"observation": observation})
    assert r.status_code == 201, r.text
    finding = r.json()
    assert finding["document_sources"], "expected the imported document to surface as context"
    assert finding["requires_human_review"] is True
    for text in (finding["observation"], finding["finding"], finding["justification"] or ""):
        assert not PLACEHOLDER_TOKEN_RE.search(text), text

    severity = "mineure" if finding["finding_type"] == "non_conformite" else None
    r = client.post(f"/sessions/{session_id}/findings/0/review", headers=headers, json={
        "approve": True, "severity": severity, "reviewer": "J. Dupont",
    })
    assert r.status_code == 200, r.text

    r = client.get(f"/sessions/{session_id}/report", headers=headers, params={"format": "docx"})
    assert r.status_code == 200
    document = docx.Document(io.BytesIO(r.content))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "Meridian" in text
    assert "SRV-API-E2E-77" in text
    assert not PLACEHOLDER_TOKEN_RE.search(text), PLACEHOLDER_TOKEN_RE.findall(text)

    r = client.post(f"/sessions/{session_id}/close", headers=headers)
    assert r.status_code == 200 and r.json()["status"] == "closed"
