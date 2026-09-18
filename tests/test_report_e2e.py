"""End-to-end Stage 4: real Qdrant + real Ollama (embedding, reranker, generation) + real corpus.

A full session built from real `AuditAssistant.analyze()` calls on observations distinct from
every other benchmark/eval/e2e set in this project, reviewed with a realistic approve/reject/edit
mix, rendered to both formats, and the rendered files checked for real content -- including that
nothing pending or rejected leaked into either document.
"""

import docx
import pypdf
import pytest

from src.audit_assistant import AuditAssistant
from src.config import get_settings
from src.report.builder import build_report
from src.report.docx_renderer import render_docx
from src.report.pdf_renderer import render_pdf
from src.report.session import AuditSession, ReviewDecision, Severity, SessionMetadata, TeamMember
from tests.conftest import huggingface_reachable, ollama_generation_model_available, ollama_model_available

pytestmark = pytest.mark.e2e

OBSERVATIONS = [
    "Chez la société Meridian, les clés de chiffrement utilisées pour protéger les sauvegardes "
    "ne sont jamais renouvelées depuis leur création initiale.",
    "L'accès à la salle serveurs n'est protégé que par une serrure à clé partagée entre tous les "
    "employés, sans traçabilité des entrées et sorties.",
    "Le distributeur de boissons du hall d'accueil n'accepte plus les paiements sans contact.",  # off-domain
]


@pytest.fixture(scope="module")
def reviewed_session():
    settings = get_settings()
    if not (ollama_model_available(settings) and ollama_generation_model_available(settings) and huggingface_reachable()):
        pytest.skip("Embedding/generation Ollama model or huggingface.co not reachable")

    assistant = AuditAssistant.from_settings(settings)
    session = AuditSession(metadata=SessionMetadata(
        title="Rapport d'audit ISO/IEC 27001 - Meridian", client_name="Meridian",
        scope="Sécurité physique et cryptographie du périmètre audité.",
        standards=["ISO/IEC 27001:2022"], audit_team=[TeamMember(name="J. Dupont", role="Lead Auditor")],
        reference="MISSION-E2E-01",
    ))
    for observation in OBSERVATIONS:
        session.add_finding(assistant, observation)

    # A realistic mix: approve the two real findings (one edited), reject the off-domain one.
    session.review(0, ReviewDecision.APPROVED, reviewer="J. Dupont", edited_text=(
        session.findings[0].finding.finding + " [confirmé par l'auditeur]"
    ))
    session.review(1, ReviewDecision.APPROVED, reviewer="J. Dupont")
    session.review(2, ReviewDecision.REJECTED, reviewer="J. Dupont")

    for i, reviewed in enumerate(session.findings):
        if reviewed.decision == ReviewDecision.APPROVED and reviewed.finding.finding_type.value == "non_conformite":
            session.review(i, ReviewDecision.APPROVED, severity=Severity.MINEURE)

    return session


def test_two_of_three_observations_are_approved_one_rejected(reviewed_session):
    assert len(reviewed_session.approved) == 2
    assert reviewed_session.rejected_count == 1


def test_report_builds_from_the_reviewed_session(reviewed_session):
    report = build_report(reviewed_session)
    assert report.title == reviewed_session.metadata.title
    assert "Meridian" in report.subtitle


class TestRenderedDocuments:
    @pytest.fixture(scope="class")
    def rendered(self, reviewed_session, tmp_path_factory):
        report = build_report(reviewed_session)
        out_dir = tmp_path_factory.mktemp("report_e2e")
        docx_path, pdf_path = out_dir / "rapport.docx", out_dir / "rapport.pdf"
        render_docx(report, docx_path)
        render_pdf(report, pdf_path)
        return docx_path, pdf_path

    def _all_text(self, path):
        if str(path).endswith(".docx"):
            document = docx.Document(str(path))
            parts = [p.text for p in document.paragraphs]
            parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            return "\n".join(parts)
        reader = pypdf.PdfReader(str(path))
        return "\n".join(page.extract_text() for page in reader.pages)

    @pytest.mark.parametrize("which", [0, 1])
    def test_real_client_name_and_reference_appear_in_both_formats(self, rendered, which):
        text = self._all_text(rendered[which])
        assert "Meridian" in text and "MISSION-E2E-01" in text

    @pytest.mark.parametrize("which", [0, 1])
    def test_approved_edited_finding_text_appears(self, rendered, which, reviewed_session):
        text = self._all_text(rendered[which])
        assert "[confirmé par l'auditeur]" in text

    @pytest.mark.parametrize("which", [0, 1])
    def test_rejected_finding_text_never_appears(self, rendered, which, reviewed_session):
        # The finding's own drafted text can be a short, generic label (observed: a bare
        # "Non-conformité", which would trivially -- and wrongly -- match the section heading
        # "Non-conformités" that legitimately exists because of the *other*, approved finding).
        # The original observation's distinctive wording is a robust, unique marker instead.
        rejected_observation = reviewed_session.findings[2].finding.observation
        assert "distributeur de boissons" in rejected_observation  # sanity: this is the off-domain one
        assert "distributeur de boissons" not in self._all_text(rendered[which])

    @pytest.mark.parametrize("which", [0, 1])
    def test_any_cited_iso_reference_appears_in_the_rendered_document(self, rendered, which, reviewed_session):
        text = self._all_text(rendered[which])
        cited = {ref for f in reviewed_session.approved for ref in f.finding.iso_reference}
        for ref in cited:
            assert ref in text
