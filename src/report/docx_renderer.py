"""Renders a Report to a .docx file.

Independent from pdf_renderer.py: both only walk the Report/ReportSection/Block model
(models.py), so neither duplicates the cahier des charges section logic that lives in builder.py.
"""

from pathlib import Path
from typing import Union

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src.report.models import BulletList, Paragraph, Report, Table


def render_docx(report: Report, path: Path) -> None:
    document = Document()
    document.add_heading(report.title, level=0)
    if report.subtitle:
        subtitle = document.add_paragraph(report.subtitle)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for section in report.sections:
        document.add_heading(section.title, level=1)
        for block in section.blocks:
            _render_block(document, block)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _render_block(document: Document, block: Union[Paragraph, BulletList, Table]) -> None:
    if isinstance(block, Paragraph):
        run = document.add_paragraph().add_run(block.text)
        run.bold = block.bold
    elif isinstance(block, BulletList):
        for item in block.items:
            document.add_paragraph(item, style="List Bullet")
    elif isinstance(block, Table):
        table = document.add_table(rows=1, cols=len(block.headers))
        table.style = "Table Grid"
        for cell, header in zip(table.rows[0].cells, block.headers):
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True
        for row in block.rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = value
    else:
        raise TypeError(f"Unknown block type: {type(block)}")
