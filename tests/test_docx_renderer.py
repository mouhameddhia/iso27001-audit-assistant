"""render_docx(): content round-trip -- reopen the written file and check the real content is
there, not just that a file was produced.
"""

import docx

from src.report.docx_renderer import render_docx
from src.report.models import BulletList, Paragraph, Report, ReportSection, Table


def all_text(document: docx.Document) -> str:
    parts = [p.text for p in document.paragraphs]
    parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def test_title_and_subtitle_are_written(tmp_path):
    path = tmp_path / "report.docx"
    render_docx(Report(title="Rapport d'audit", subtitle="Contoso — MISSION-001", sections=[]), path)

    document = docx.Document(str(path))
    assert document.paragraphs[0].text == "Rapport d'audit"
    assert "Contoso — MISSION-001" in all_text(document)


def test_section_titles_become_headings(tmp_path):
    path = tmp_path / "report.docx"
    report = Report("T", "", [ReportSection("Constats", [Paragraph("Un constat.")])])
    render_docx(report, path)

    document = docx.Document(str(path))
    heading = next(p for p in document.paragraphs if p.text == "Constats")
    assert heading.style.name.startswith("Heading")


def test_bold_paragraph_is_rendered_bold(tmp_path):
    path = tmp_path / "report.docx"
    report = Report("T", "", [ReportSection("S", [Paragraph("NC-001 (Majeure)", bold=True), Paragraph("Corps normal.")])])
    render_docx(report, path)

    document = docx.Document(str(path))
    header = next(p for p in document.paragraphs if p.text == "NC-001 (Majeure)")
    body = next(p for p in document.paragraphs if p.text == "Corps normal.")
    assert header.runs[0].bold is True
    assert not body.runs[0].bold


def test_bullet_list_items_are_all_present(tmp_path):
    path = tmp_path / "report.docx"
    report = Report("T", "", [ReportSection("Conclusion", [BulletList(["ISO/IEC 27001 A.5.18", "ISO/IEC 27001 A.8.8"])])])
    render_docx(report, path)

    text = all_text(docx.Document(str(path)))
    assert "ISO/IEC 27001 A.5.18" in text and "ISO/IEC 27001 A.8.8" in text


def test_table_headers_and_rows_are_present(tmp_path):
    path = tmp_path / "report.docx"
    report = Report("T", "", [ReportSection("Informations générales", [
        Table(["Champ", "Valeur"], [["Client", "Contoso"], ["Référence", "MISSION-001"]])
    ])])
    render_docx(report, path)

    document = docx.Document(str(path))
    assert len(document.tables) == 1
    rows = [[c.text for c in row.cells] for row in document.tables[0].rows]
    assert rows == [["Champ", "Valeur"], ["Client", "Contoso"], ["Référence", "MISSION-001"]]


def test_unicode_and_markup_like_characters_survive_intact(tmp_path):
    path = tmp_path / "report.docx"
    text = "Écart : accès & privilèges < seuil, à revoir."
    render_docx(Report("T", "", [ReportSection("S", [Paragraph(text)])]), path)
    assert text in all_text(docx.Document(str(path)))


def test_creates_parent_directories(tmp_path):
    path = tmp_path / "nested" / "dir" / "report.docx"
    render_docx(Report("T", "", []), path)
    assert path.exists()
