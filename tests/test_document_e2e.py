"""End-to-end document import: real Qdrant + real Ollama (embedding, generation) + real corpus.

The exact scenario the Document Import plan was designed around: import a synthetic "previous
audit report" containing a company name with *no* anonymization-triggering keyword nearby (caught
only via `client_aliases`) and a server name that *does* have one, confirm it, draft a finding
whose observation references the same entities, and check end to end that:
  - the sidecar file a human would review never contains the real values,
  - using client-document context forces human review,
  - the final report restores the real values and never leaks a raw placeholder token.
"""

import re

import docx
import pypdf
import pytest

from src.audit_assistant import AuditAssistant
from src.config import get_settings
from src.documents.ingestion import close_session_collection, confirm_import, propose_import, session_collection_name
from src.documents.retrieval import build_document_retriever
from src.report.builder import build_report
from src.report.docx_renderer import render_docx
from src.report.pdf_renderer import render_pdf
from src.report.session import AuditSession, ReviewDecision, Severity, SessionMetadata, TeamMember
from src.vectorstore import QdrantStore
from tests.conftest import huggingface_reachable, ollama_generation_model_available, ollama_model_available, qdrant_reachable

pytestmark = pytest.mark.e2e

PLACEHOLDER_TOKEN_RE = re.compile(r"\b[A-Z]+_\d{3}\b")

DOCUMENT_PARAGRAPHS = [
    "Rapport d'audit precedent.",
    "Le present rapport concerne l'audit de securite realise pour Meridian en 2024.",
    "Le serveur SRV-LEGACY-42 conservait des journaux d'audit non chiffres depuis sa mise en "
    "service sans que cela ait ete corrige a la suite de la precedente mission.",
]

OBSERVATION = (
    "Chez Meridian, le serveur SRV-LEGACY-42 conserve encore des journaux d'audit non chiffres, "
    "comme deja releve dans un rapport precedent."
)


def _all_text(path) -> str:
    if str(path).endswith(".docx"):
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        parts += [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        return "\n".join(parts)
    reader = pypdf.PdfReader(str(path))
    return "\n".join(page.extract_text() for page in reader.pages)


@pytest.fixture(scope="module")
def settings():
    settings = get_settings()
    if not (
        qdrant_reachable(settings) and ollama_model_available(settings)
        and ollama_generation_model_available(settings) and huggingface_reachable()
    ):
        pytest.skip("Qdrant, embedding/generation Ollama model, or huggingface.co not reachable")
    return settings


@pytest.fixture(scope="module")
def imported_session(settings, tmp_path_factory):
    session = AuditSession(metadata=SessionMetadata(
        title="Rapport d'audit ISO/IEC 27001 - Meridian Holdings", client_name="Meridian Holdings",
        client_aliases=["Meridian"], scope="Perimetre de test document import.",
        standards=["ISO/IEC 27001:2022"], audit_team=[TeamMember(name="J. Dupont", role="Lead Auditor")],
        reference="MISSION-DOC-E2E-01",
    ))
    settings = settings.model_copy(update={"document_import_dir": tmp_path_factory.mktemp("imports")})

    import docx as docx_module
    source = tmp_path_factory.mktemp("source") / "rapport_precedent.docx"
    document = docx_module.Document()
    for paragraph in DOCUMENT_PARAGRAPHS:
        document.add_paragraph(paragraph)
    document.save(str(source))

    record = propose_import(session, source, doc_type="previous_report", settings=settings)
    yield session, settings, record

    store = QdrantStore.from_settings(settings, collection=session_collection_name(session, settings))
    close_session_collection(session, settings=settings, store=store)


def test_the_reviewable_sidecar_never_contains_the_real_client_or_server_name(imported_session):
    session, settings, record = imported_session
    from pathlib import Path
    sidecar_text = Path(record.sidecar_path).read_text(encoding="utf-8")
    assert "Meridian" not in sidecar_text
    assert "SRV-LEGACY-42" not in sidecar_text


@pytest.fixture(scope="module")
def confirmed_session(imported_session):
    session, settings, record = imported_session
    confirm_import(session, record.doc_id, settings=settings)
    return session, settings


def test_confirm_import_stores_chunks_and_marks_confirmed(confirmed_session):
    session, settings = confirmed_session
    assert session.confirmed_documents
    assert session.confirmed_documents[0].chunk_count > 0


@pytest.fixture(scope="module")
def finding_with_document_context(confirmed_session):
    session, settings = confirmed_session
    anonymizer = session.anonymizer
    assistant = AuditAssistant.from_settings(settings, anonymizer=anonymizer)
    document_retriever = build_document_retriever(session, settings)
    assert document_retriever is not None

    reviewed = session.add_finding(assistant, OBSERVATION, document_retriever=document_retriever)
    session.sync_anonymizer(anonymizer)
    return session, reviewed


def test_document_context_was_actually_retrieved_and_used(finding_with_document_context):
    _, reviewed = finding_with_document_context
    assert reviewed.finding.document_sources, "expected the imported document to surface as context"


def test_using_document_context_forces_human_review(finding_with_document_context):
    _, reviewed = finding_with_document_context
    assert reviewed.finding.requires_human_review is True


def test_no_raw_placeholder_ever_reaches_the_finding_after_deanonymization(finding_with_document_context):
    _, reviewed = finding_with_document_context
    for text in (reviewed.finding.observation, reviewed.finding.finding, reviewed.finding.justification or ""):
        assert not PLACEHOLDER_TOKEN_RE.search(text), text


class TestFinalReport:
    @pytest.fixture(scope="class")
    def rendered(self, finding_with_document_context, tmp_path_factory):
        session, reviewed = finding_with_document_context
        index = session.findings.index(reviewed)
        severity = Severity.MINEURE if reviewed.finding.finding_type.value == "non_conformite" else None
        session.review(index, ReviewDecision.APPROVED, reviewer="J. Dupont", severity=severity)

        report = build_report(session)
        out_dir = tmp_path_factory.mktemp("document_e2e_report")
        docx_path, pdf_path = out_dir / "rapport.docx", out_dir / "rapport.pdf"
        render_docx(report, docx_path)
        render_pdf(report, pdf_path)
        return docx_path, pdf_path

    @pytest.mark.parametrize("which", [0, 1])
    def test_real_client_and_server_name_are_restored_in_the_report(self, rendered, which):
        text = _all_text(rendered[which])
        assert "Meridian" in text
        assert "SRV-LEGACY-42" in text

    @pytest.mark.parametrize("which", [0, 1])
    def test_no_placeholder_token_ever_appears_in_the_rendered_report(self, rendered, which):
        text = _all_text(rendered[which])
        assert not PLACEHOLDER_TOKEN_RE.search(text), PLACEHOLDER_TOKEN_RE.findall(text)
